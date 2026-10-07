"""Deterministic, same-origin SG-to-REGCP1/PLL2/REECB1 model construction."""
from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

import openpyxl

BUILD_VERSION = "1.0"
COMMON_SHEETS = ("Bus", "PQ", "PV", "Slack", "Shunt", "Line", "Area", "BusFreq", "ACEc", "Toggler")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_books(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    books = {}
    for sheet in workbook:
        rows = list(sheet.values)
        books[sheet.title] = [{key: value for key, value in zip(rows[0], row)} for row in rows[1:]
                             if any(value is not None for value in row)]
    workbook.close()
    return books


def write_deterministic_xlsx(path, books):
    """Normalize workbook properties and ZIP metadata, as well as model values."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "ieee39ts"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    for name, rows in books.items():
        if not rows:
            continue
        sheet = workbook.create_sheet(name)
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(key) for key in headers])
    buffer = BytesIO()
    workbook.save(buffer)
    normalized = BytesIO()
    with ZipFile(buffer) as source, ZipFile(normalized, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+", rb"\g<1>2000-01-01T00:00:00Z", content)
            entry = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.create_system = 3
            entry.external_attr = 0o600 << 16
            target.writestr(entry, content, compresslevel=9)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(normalized.getvalue())


def renewable_share(static_generation, renewable_sources):
    """Count post-PFlow online StaticGen injections once; never add dynamic Pe."""
    online = {row["static_gen_id"]: row for row in static_generation if row["online"]}
    if len(online) != sum(bool(row["online"]) for row in static_generation):
        raise ValueError("Duplicate online StaticGen identity")
    represented = [row["gen"] for row in renewable_sources if row.get("u", 1)]
    if len(set(represented)) != len(represented):
        raise ValueError("A fully replaced StaticGen must have one renewable dynamic source")
    if not set(represented).issubset(online):
        raise ValueError("Renewable source points to missing/offline StaticGen")
    total = sum(row["P0_MW"] for row in online.values())
    if total <= 0:
        raise ValueError("Total online generation must be positive")
    renewable = sum(online[idx]["P0_MW"] for idx in represented)
    return {"renewable_P_MW": renewable, "total_generation_P_MW": total,
            "renewable_share_pct": 100 * renewable / total}


def replace_generators(baseline, pilot, buses):
    buses = tuple(buses)
    if len(set(buses)) != len(buses):
        raise ValueError("Replacement buses must be unique")
    if set(buses) & {row["bus"] for row in baseline["Slack"] if row["u"]}:
        raise ValueError("Cannot replace the Slack synchronous generator")
    original = {row["bus"]: row for row in baseline["GENROU"] if row["u"]}
    if not set(buses).issubset(original):
        raise ValueError("Replacement bus has no online GENROU")
    books = deepcopy(baseline)
    removed = {original[bus]["idx"] for bus in buses}
    removal = {"GENROU": sorted(removed)}
    # Follow the actual syn and avr foreign keys, including exciter-linked PSS.
    for sheet in ("GENROU", "TGOV1N", "IEEEX1", "IEEEST"):
        discarded = [row for row in books[sheet] if row["idx"] in removed
                     or any(row.get(key) in removed for key in ("syn", "avr"))]
        removed.update(row["idx"] for row in discarded)
        removal[sheet] = [row["idx"] for row in discarded]
        books[sheet] = [row for row in books[sheet] if row not in discarded]
        for uid, row in enumerate(books[sheet]):
            row["uid"] = uid
    replacements = []
    for sheet in ("PLL2", "REGCP1", "REECB1"):
        books[sheet] = []
    pilot_sn = float(pilot["REGCP1"][0]["Sn"])
    for uid, bus in enumerate(buses):
        synch = original[bus]
        sn = float(synch["Sn"])
        reg_id, pll_id, reec_id = f"REGCP1_B{bus}", f"PLL2_B{bus}", f"REECB1_B{bus}"
        pll = deepcopy(pilot["PLL2"][0])
        pll.update(uid=uid, idx=pll_id, name=pll_id, bus=bus)
        reg = deepcopy(pilot["REGCP1"][0])
        reg.update(uid=uid, idx=reg_id, name=reg_id, bus=bus, gen=synch["gen"], Sn=sn,
                   pll=pll_id, gammap=1., gammaq=1.)
        reec = deepcopy(pilot["REECB1"][0])
        reec.update(uid=uid, idx=reec_id, name=reec_id, reg=reg_id)
        # ANDES 2.0.0 does not automatically base-convert these three inputs.
        # Preserve the pilot's normalized device-base response at each rating.
        for key in ("Kqv", "Iqh1", "Iql1"):
            reec[key] = float(reec[key]) * sn / pilot_sn
        books["PLL2"].append(pll)
        books["REGCP1"].append(reg)
        books["REECB1"].append(reec)
        replacements.append({"bus": bus, "generator_id": synch["idx"], "static_gen_id": synch["gen"],
                             "original_synch_Sn_MVA": sn, "renewable_Sn_MVA": sn,
                             "replacement_rule": "Equal Sn; retain StaticGen and all its parameters; gammaP=gammaQ=1",
                             "removed": {
                                 "GENROU": [synch["idx"]],
                                 "TGOV1N": [r["idx"] for r in baseline["TGOV1N"] if r["syn"] == synch["idx"]],
                                 "IEEEX1": [r["idx"] for r in baseline["IEEEX1"] if r["syn"] == synch["idx"]],
                                 "IEEEST": [r["idx"] for r in baseline["IEEEST"] if r["avr"] in {
                                     exc["idx"] for exc in baseline["IEEEX1"] if exc["syn"] == synch["idx"]}]}})
    for name in COMMON_SHEETS:
        if books[name] != baseline[name]:
            raise ValueError(f"Shared network/static data changed: {name}")
    return books, replacements


def build_models(root: Path, config: dict, output: Path, *, smoke_report=None):
    """Rebuild model bytes and manifest without running power flow or TDS."""
    root, output = Path(root), Path(output)
    adapted_path, pilot_path = root / config["adapted_source"], root / config["pilot_source"]
    baseline, pilot = read_books(adapted_path), read_books(pilot_path)
    order = config["replacement_order"]
    if len(order) != 3 or order[0] != 37 or config["renewable_Sn_rule"] != "equal_to_original_synchronous_Sn":
        raise ValueError("Require three distinct replacements, starting at Bus37, with equal Sn")
    # Validate the entire ordered replacement plan before writing any model.
    replace_generators(baseline, pilot, order)
    output.mkdir(parents=True, exist_ok=True)
    script_path = root / "scripts/build_renewable_scenarios.py"
    manifest = {"schema_version": 1, "build_version": BUILD_VERSION,
                "sources": [{"file": config[key], "sha256": sha256(root / config[key])}
                            for key in ("adapted_source", "pilot_source")],
                "generator": {"script": "scripts/build_renewable_scenarios.py", "sha256": sha256(script_path),
                              "module_sha256": sha256(Path(__file__)), "openpyxl_version": openpyxl.__version__},
                "configuration": config,
                "common_sheet_fingerprint": hashlib.sha256(json.dumps({k:baseline[k] for k in COMMON_SHEETS},
                                                        sort_keys=True, allow_nan=False).encode()).hexdigest(),
                "template_rule": {"gammap": 1, "gammaq": 1, "renewable_Sn": "original GENROU.Sn",
                                  "static_gen_parameters": "unchanged",
                                  "capacity_scaling": "Kqv/Iqh1/Iql1 multiplied by original Sn / pilot Sn; other template inputs unchanged"},
                "pilot_crosscheck": {"pilot_bus37_Sn_MVA": float(pilot['REGCP1'][0]['Sn']),
                                     "formal_bus37_Sn_MVA": float(next(row['Sn'] for row in baseline['GENROU'] if row['bus']==37)),
                                     "intentional_differences": ["Keep StaticGen.Sn from S0", "Match renewable Sn to original SG", "Scale three unconverted controller inputs", "Explicit gammaP/gammaQ=1"]},
                "scenarios": []}
    for index in range(4):
        buses = order[:index]
        filename = "S0_adapted_sg.xlsx" if not buses else f"S{index}_" + "_".join(f"bus{bus}" for bus in buses) + "_renewable.xlsx"
        if not buses:
            (output / filename).write_bytes(adapted_path.read_bytes())
            replacements = []
        else:
            books, replacements = replace_generators(baseline, pilot, buses)
            write_deterministic_xlsx(output / filename, books)
        model_hash = sha256(output / filename)
        scenario = {"scenario_id": f"S{index}", "file": filename, "sha256": model_hash,
                    "source_model": config["adapted_source"], "renewable_buses": buses,
                    "num_synch_generators": 10-index, "num_renewable_generators": index,
                    "replacements": replacements, "renewable_P_MW": None,
                    "total_generation_P_MW": None, "renewable_share_pct": None,
                    "share_basis": "PENDING individual scenario PFlow"}
        if smoke_report is not None:
            accepted = next(row for row in smoke_report["scenarios"] if row["scenario_id"] == f"S{index}")
            if accepted["model_sha256"] != model_hash:
                raise ValueError("Smoke report model hash differs from deterministic generated model")
            scenario.update({key: accepted[key] for key in ("renewable_P_MW", "total_generation_P_MW", "renewable_share_pct")})
            scenario["share_basis"] = "Own online StaticGen actual P after PFlow, captured before TDS.init"
            scenario["smoke_passed"] = accepted["smoke_passed"]
        manifest["scenarios"].append(scenario)
    return manifest
